# ruff: noqa: RUF001
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.modules.price_lists.domain.errors import PriceListError, PriceListErrorCode
from app.modules.price_lists.domain.mapping import CanonicalField, ColumnMapping, detect_columns
from app.modules.price_lists.infrastructure.xlsx_extractor import XlsxPriceListExtractor


def workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_aliases_are_deterministic_and_duplicate_candidates_are_ambiguous() -> None:
    columns, ambiguous = detect_columns([" قيمت (ريال) ", "نام  كالا", "کد کالا", "SKU"])
    assert columns == {CanonicalField.PRICE_IRR: 1, CanonicalField.PRODUCT_NAME: 2}
    assert ambiguous == {CanonicalField.PRODUCT_CODE: (3, 4)}


@pytest.mark.parametrize(
    "columns,code",
    [
        ({CanonicalField.PRODUCT_CODE: 1}, "REQUIRED_COLUMN_MISSING"),
        (
            {
                CanonicalField.PRODUCT_CODE: 1,
                CanonicalField.PRODUCT_NAME: 1,
                CanonicalField.PRICE_IRR: 3,
            },
            "INVALID_COLUMN_MAPPING",
        ),
        (
            {
                CanonicalField.PRODUCT_CODE: 1,
                CanonicalField.PRODUCT_NAME: 2,
                CanonicalField.PRICE_IRR: 4,
            },
            "INVALID_COLUMN_MAPPING",
        ),
    ],
)
def test_mapping_validation_blocks_missing_duplicate_and_out_of_range(columns, code) -> None:
    with pytest.raises(PriceListError) as raised:
        ColumnMapping("Sheet", 1, columns).validate(column_count=3)
    assert raised.value.code.value == code


def test_ambiguous_header_never_extracts_successfully() -> None:
    content = workbook_bytes(
        [["کد کالا", "SKU", "نام کالا", "قیمت ریال"], ["001", "002", "کابل", 1000]]
    )
    report = XlsxPriceListExtractor().analyze(content, filename="test.xlsx")
    assert report.error.code is PriceListErrorCode.AMBIGUOUS_COLUMN_MAPPING
    assert report.items == []
    assert report.snapshot["preview"][0]["source_row_number"] == 2


def test_explicit_header_and_mapping_preserve_source_and_leading_zeroes() -> None:
    content = workbook_bytes([["فهرست"], ["مبلغ", "شرح", "شناسه"], ["۱٬۲۳۴", "كابل", "۰۰۱"]])
    mapping = ColumnMapping(
        "Sheet",
        2,
        {
            CanonicalField.PRODUCT_CODE: 3,
            CanonicalField.PRODUCT_NAME: 2,
            CanonicalField.PRICE_IRR: 1,
        },
    )
    report = XlsxPriceListExtractor().analyze(content, filename="test.xlsx", mapping=mapping)
    assert report.error is None
    assert report.items[0].source_row_number == 3
    assert report.items[0].product_code_normalized == "001"
    assert report.items[0].price_irr == 1234


def test_row_errors_collect_numbers_fields_and_keep_no_partial_rows() -> None:
    content = workbook_bytes(
        [
            ["کد کالا", "نام کالا", "قیمت ریال"],
            ["001", "کابل", 1000],
            ["", "کابل", 1000],
            ["002", "کابل", "12.5"],
            ["001", "کابل", 1000],
        ]
    )
    report = XlsxPriceListExtractor().analyze(content, filename="test.xlsx")
    assert report.items == []
    assert [
        (issue["row"], issue["field"], issue["code"]) for issue in report.snapshot["issues"]
    ] == [
        (3, "product_code", "EMPTY_PRODUCT_CODE"),
        (4, "price_irr", "INVALID_PRICE"),
        (5, "product_code", "DUPLICATE_PRODUCT_CODE"),
    ]


def test_template_with_missing_sheet_is_failed_and_keeps_reviewable_proposal() -> None:
    content = workbook_bytes([["کد کالا", "نام کالا", "قیمت ریال"], ["001", "کابل", 1000]])
    report = XlsxPriceListExtractor().analyze(
        content,
        filename="prices.xlsx",
        mapping=ColumnMapping(
            "previous",
            1,
            {
                CanonicalField.PRODUCT_CODE: 1,
                CanonicalField.PRODUCT_NAME: 2,
                CanonicalField.PRICE_IRR: 3,
            },
        ),
        header_signature=("کدکالا", "نامکالا", "قیمتریال"),
    )
    assert report.error.code is PriceListErrorCode.TEMPLATE_LAYOUT_MISMATCH
    assert report.snapshot["preview"]
    assert report.items == []


def test_renamed_macro_enabled_content_is_rejected() -> None:
    from zipfile import ZipFile

    content = workbook_bytes([["کد کالا", "نام کالا", "قیمت ریال"], ["001", "کابل", 1000]])
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(name, data)
    with pytest.raises(PriceListError) as raised:
        XlsxPriceListExtractor().extract(output.getvalue(), filename="renamed.xlsx")
    assert raised.value.code is PriceListErrorCode.INVALID_FILE_TYPE


@pytest.mark.parametrize("value", [1000.0, "1,2", "1,,000", "12٬34", "1 2"])
def test_ambiguous_prices_are_rejected(value: object) -> None:
    from app.modules.price_lists.domain.normalization import parse_price_irr

    with pytest.raises(PriceListError) as raised:
        parse_price_irr(value)
    assert raised.value.code is PriceListErrorCode.INVALID_PRICE


def test_header_alias_order_and_preamble_detection() -> None:
    report = XlsxPriceListExtractor().analyze(
        workbook_bytes(
            [
                ["فهرست"],
                ["تاریخ"],
                ["قیمت (ریال)", " نام  كالا ", "ITEM CODE"],
                ["۱۲٬۳۴۵", "کابل", "00-ab"],
            ]
        ),
        filename="prices.xlsx",
    )
    assert report.error is None
    assert report.snapshot["used_mapping"]["header_row"] == 3
    assert report.items[0].price_irr == 12345
    assert report.items[0].product_code_normalized == "00-AB"


def test_ambiguous_header_and_multiple_sheets_require_explicit_choice() -> None:
    workbook = Workbook()
    for sheet in [workbook.active, workbook.create_sheet("Second")]:
        sheet.append(["کد کالا", "نام کالا", "قیمت ریال"])
        sheet.append(["001", "کابل", 1000])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    report = XlsxPriceListExtractor().analyze(output.getvalue(), filename="prices.xlsx")
    assert report.error.code is PriceListErrorCode.AMBIGUOUS_HEADER
    assert report.items == []


def test_errors_and_preview_are_bounded_and_raw_duplicate_headers_are_preserved() -> None:
    rows = [["کد کالا", "نام کالا", "قیمت ریال", "شرح", "شرح"]] + [
        ["", "کابل", "invalid", "<script>" + ("x" * 500), "other"] for _ in range(110)
    ]
    report = XlsxPriceListExtractor().analyze(workbook_bytes(rows), filename="prices.xlsx")
    assert len(report.snapshot["issues"]) == 100
    assert report.snapshot["total_errors"] == 110
    assert report.snapshot["errors_truncated"] is True
    assert len(report.snapshot["preview"]) == 5
    assert len(report.snapshot["preview"][0]["values"][3]) == 200
    valid = XlsxPriceListExtractor().extract(
        workbook_bytes([rows[0], ["001", "کابل", 1000, "first", "second"]]), filename="prices.xlsx"
    )[0]
    assert valid.raw_row["column_4"] == "first"
    assert valid.raw_row["column_5"] == "second"


def test_manual_mapping_cannot_bypass_toman_preamble() -> None:
    mapping = ColumnMapping(
        "Sheet",
        2,
        {
            CanonicalField.PRODUCT_CODE: 1,
            CanonicalField.PRODUCT_NAME: 2,
            CanonicalField.PRICE_IRR: 3,
        },
    )
    report = XlsxPriceListExtractor().analyze(
        workbook_bytes([["مبالغ به تومان"], ["کد", "نام", "نرخ"], ["001", "کابل", 1000]]),
        filename="prices.xlsx",
        mapping=mapping,
    )
    assert report.error.code is PriceListErrorCode.TOMAN_NOT_SUPPORTED
    assert report.items == []


def test_processing_deadline_is_enforced_before_persistence() -> None:
    from app.modules.price_lists.infrastructure.xlsx_extractor import XlsxExtractionLimits

    with pytest.raises(PriceListError) as raised:
        XlsxPriceListExtractor(XlsxExtractionLimits(max_processing_seconds=0.000001)).analyze(
            workbook_bytes([["کد کالا", "نام کالا", "قیمت ریال"], ["001", "کابل", 1000]]),
            filename="prices.xlsx",
        )
    assert raised.value.code is PriceListErrorCode.WORKBOOK_PROCESSING_TIMEOUT


def test_malformed_workbook_xml_is_a_stable_file_error() -> None:
    from zipfile import ZipFile

    content = workbook_bytes([["کد کالا", "نام کالا", "قیمت ریال"]])
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            target.writestr(name, b"<broken" if name == "xl/workbook.xml" else source.read(name))
    with pytest.raises(PriceListError) as raised:
        XlsxPriceListExtractor().analyze(output.getvalue(), filename="broken.xlsx")
    assert raised.value.code is PriceListErrorCode.WORKBOOK_UNREADABLE


def test_underreported_sheet_dimensions_cannot_hide_an_ambiguous_price_column() -> None:
    import re
    from zipfile import ZipFile

    content = workbook_bytes(
        [["کد کالا", "نام کالا", "قیمت ریال", "قیمت خرید ریال"], ["001", "کابل", 1000, 2000]]
    )
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:C2"', data)
            target.writestr(name, data)
    report = XlsxPriceListExtractor().analyze(output.getvalue(), filename="prices.xlsx")
    assert report.error.code is PriceListErrorCode.AMBIGUOUS_COLUMN_MAPPING
    assert report.items == []
